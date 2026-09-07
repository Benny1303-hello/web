import json
import re
import sys
from pathlib import Path

import openpyxl

SRC = r"Z:\Data\Lam\Copy of BẢNG GIÁ EATON.xlsx"
OUT = Path(__file__).resolve().parents[4] / "data" / "product-imports" / "copy-of-bang-gia-eaton.json"

wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb["Sheet1"]

title = ws["A1"].value
header = [c.value for c in ws[2]]
assert header == ["Mã hàng", "Catalogue Number ", "Giá MD full v", "BH", "XS", "Tên xuất hóa đơn", "Giá TTC chào"], header

category = title.replace("BẢNG GIÁ", "").strip() if title else "UPS EATON"

VA_W_RE = re.compile(r"(\d{2,6})\s*VA\s*/\s*(\d{2,6})\s*W", re.IGNORECASE)
VOLT_RE = re.compile(r"\b(\d{3}(?:/\d{3}){0,2})\s*V\b")

products = []
skipped_blank = 0
issues = []
seen_part_numbers = {}

for row in ws.iter_rows(min_row=3, max_row=ws.max_row):
    vals = [c.value for c in row]
    r = row[0].row
    if all(v is None for v in vals):
        skipped_blank += 1
        continue

    ma_hang, catalogue_number, price_md, warranty, stock, name, price_ttc = vals

    part_number = catalogue_number or ma_hang
    if not part_number:
        issues.append({"row": r, "issue": "no code in either Mã hàng or Catalogue Number", "raw": vals})

    specs = {}
    if name:
        m = VA_W_RE.search(name)
        if m:
            specs["capacity_va"] = int(m.group(1))
            specs["capacity_w"] = int(m.group(2))
        v = VOLT_RE.search(name)
        if v:
            specs["voltage"] = v.group(1) + "V"
    else:
        issues.append({"row": r, "issue": "missing Tên xuất hóa đơn (name)", "raw": vals})

    if not warranty:
        issues.append({"row": r, "issue": "missing BH (warranty)", "raw": vals})

    if catalogue_number is None:
        issues.append({"row": r, "issue": "missing Catalogue Number, fell back to Mã hàng as part_number", "raw": vals})

    prices = []
    if price_md is not None:
        prices.append({"label": "Giá MD full v", "value": price_md})
    if price_ttc is not None:
        prices.append({"label": "Giá TTC chào", "value": price_ttc})
    if price_ttc == 0:
        issues.append({"row": r, "issue": "Giá TTC chào is 0 — likely price not set yet, not a real zero price", "raw": vals})
    if not prices:
        issues.append({"row": r, "issue": "no price found in either price column", "raw": vals})

    extra_fields = []
    if ma_hang and ma_hang != part_number:
        extra_fields.append({"label": "Mã hàng (internal code)", "value": ma_hang})

    product = {
        "row": r,
        "part_number": part_number,
        "name": name,
        "spec_text": name or "",
        "specs": specs,
        "prices": prices,
        "origin": None,
        "stock": stock,
        "warranty": warranty,
        "extra_fields": extra_fields,
    }
    products.append(product)

    if part_number:
        seen_part_numbers.setdefault(part_number, []).append(r)

duplicates = [
    {"part_number": pn, "rows": rows, "occurrences": len(rows)}
    for pn, rows in seen_part_numbers.items()
    if len(rows) > 1
]

total_rows = ws.max_row - 2  # excluding title + header rows
reconciled = len(products) + skipped_blank == total_rows

out = {
    "source_file": "Copy of BẢNG GIÁ EATON.xlsx",
    "extracted_at": "2026-09-07",
    "categories": [
        {
            "category": category,
            "source": "sheet:Sheet1",
            "products": products,
        }
    ],
    "summary": {
        "total_categories": 1,
        "total_products": len(products),
        "by_category": {category: len(products)},
        "duplicate_part_numbers": duplicates,
        "rows_with_issues": issues,
        "row_reconciliation": {
            "total_data_rows": total_rows,
            "products_found": len(products),
            "skipped_blank_rows": skipped_blank,
            "reconciled": reconciled,
        },
    },
}

OUT.parent.mkdir(parents=True, exist_ok=True)
OUT.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")

print(f"Wrote {OUT}")
print(f"Products: {len(products)}  Duplicates: {len(duplicates)}  Issues: {len(issues)}  Reconciled: {reconciled}")
for iss in issues:
    print(f"  ROW {iss['row']}: {iss['issue']}")
